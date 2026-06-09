"""Build the full Ibrahim Malik session script — print-ready, with note space."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ───── colors ─────
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GOLD = RGBColor(0xC9, 0xA9, 0x61)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = RGBColor(0x88, 0x88, 0x88)
RED = RGBColor(0xB0, 0x00, 0x00)
GREEN = RGBColor(0x1A, 0x73, 0x1A)
BLACK = RGBColor(0x00, 0x00, 0x00)

doc = Document()

# Page margins
section = doc.sections[0]
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)

# ───── helpers ─────
def shade_paragraph(paragraph, color_hex):
    """Add background shading to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    pPr.append(shd)

def add_border_bottom(paragraph, color_hex="888888", size=4):
    """Add a bottom border to a paragraph for note-taking lines."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_title(text, size=22, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Georgia"

def add_h1(text, size=16, color=NAVY, shade=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if shade:
        shade_paragraph(p, "F4ECD8")
    run = p.add_run("  " + text + "  ")
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Georgia"

def add_h2(text, size=12, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color

def add_para(text, size=11, color=BLACK, italic=False, bold=False, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    run.bold = bold
    return p

def add_script(text, size=11):
    """Add a verbatim script line — boxed/shaded so Gordon can read directly."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    shade_paragraph(p, "EEF3F8")
    run = p.add_run('"' + text + '"')
    run.font.size = Pt(size)
    run.italic = True
    run.font.color.rgb = NAVY

def add_stage_direction(text, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run("[ " + text + " ]")
    run.font.size = Pt(size)
    run.italic = True
    run.font.color.rgb = LIGHT_GRAY

def add_bullet(text, size=11, color=BLACK, bold=False):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold

def add_note_lines(count=3, label=None):
    """Add N blank lines with underlines for live note-taking."""
    if label:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("→ " + label)
        run.font.size = Pt(10)
        run.italic = True
        run.font.color.rgb = GRAY
    for _ in range(count):
        p = doc.add_paragraph(" ")
        p.paragraph_format.space_after = Pt(4)
        add_border_bottom(p)

def add_hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("─" * 65)
    run.font.color.rgb = GOLD
    run.font.size = Pt(9)

def add_warning(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    shade_paragraph(p, "FFEEEE")
    run = p.add_run("⚠ " + text)
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = RED

def page_break():
    doc.add_page_break()

# ═════════════════════════════════════════════════════════════
# COVER
# ═════════════════════════════════════════════════════════════

add_title("MIZAN × IBRAHIM MALIK")
add_title("Session Script", size=16, color=GOLD)
add_para("", size=8)
add_para("Thursday, June 4, 2026   ·   5:30 – 6:00pm Hong Kong time", size=12, color=GRAY)
add_para("Channel: Zoom (link sent)   ·   Recording: LOCAL (your laptop)", size=11, color=GRAY)
add_para("First inbound from CUHK MSA route. First real-user feedback session.", size=10, color=GOLD, italic=True)

add_hr()

add_h2("HOW TO USE THIS DOCUMENT", size=11)
add_para("PRINT IT. Have it paper-side-by-laptop during the call. Read the SCRIPT BOXES verbatim — they're calibrated. Write notes in the lined spaces. Don't improvise — this is the single highest-leverage 30 minutes in Mizan's first 30 days.",
         size=10, italic=True, color=GRAY)

# ═════════════════════════════════════════════════════════════
# PRE-SESSION CHECKLIST (5 min before)
# ═════════════════════════════════════════════════════════════
add_h1("PRE-SESSION CHECKLIST — 5 min before 5:30pm HK")

add_h2("Tech")
add_bullet("Zoom open, mic + camera tested")
add_bullet("Screen-share enabled (HE will share his screen)")
add_bullet("Local recording: Zoom → Settings → Recording → 'Record to local computer' ON")
add_bullet("Phone on silent, face down")
add_bullet("Every other app closed (Slack, email, notifications off)")

add_h2("Materials")
add_bullet("This document printed OR open on second screen")
add_bullet("Pen + notepad")
add_bullet("Mizan open in a tab (mizaninvesting.com) — for your reference only, don't show him")
add_bullet("Water within reach")

add_h2("Mindset")
add_bullet("You're a learner today, not a salesperson")
add_bullet("Silence is data — your job is to listen, not to explain")
add_bullet("If criticized: nod, write it down, don't defend")

page_break()

# ═════════════════════════════════════════════════════════════
# MIZAN CHEAT SHEET
# ═════════════════════════════════════════════════════════════
add_h1("MIZAN CHEAT SHEET — read once before the call")

add_h2("Mizan in ONE sentence")
add_para("A multi-scholar halal stock screener — six scholar profiles, every threshold cited to a primary source, every verdict reconstructable from the published rules.",
         italic=True, color=NAVY)

add_h2("The 6 scholar profiles — memorize the differentiators")

add_para("AAOIFI Standard (DEFAULT)", bold=True)
add_bullet("Gulf standard, used by most Islamic banks")
add_bullet("Debt ≤ 30% mkt cap · Cash+securities ≤ 30% · Non-permissible income ≤ 5%")
add_bullet("Purification: dividends only. Source: AAOIFI Standard No. 21")

add_para("Mufti Taqi Usmani (stricter Hanafi)", bold=True)
add_bullet("Same as AAOIFI + illiquid assets ≥ 20% of total assets")
add_bullet("Purification: dividends AND capital gains")
add_bullet("MIZAN'S BIGGEST DIFFERENTIATOR. No other halal app implements this", color=GOLD, bold=True)

add_para("DJIM (Dow Jones Islamic Market)", bold=True)
add_bullet("Most permissive of major indices")
add_bullet("Debt ≤ 33% · Cash ≤ 33% · Receivables ≤ 33% · Non-perm income ≤ 5%")
add_bullet("Purification: dividends only")

add_para("Ja'fari (Shia, Sistani-aligned)", bold=True)
add_bullet("Shia, follows Ayatollah Sistani's rulings")
add_bullet("Debt ≤ 33% · Cash+securities ≤ 33%")
add_bullet("Purification: khums (20%) on net annual surplus")

add_para("Strict / Zero-tolerance", bold=True)
add_bullet("Minority view: any riba is impermissible. 0% across the board")
add_bullet("Almost no public equities pass — that's the point")

add_para("Custom (user sets own thresholds)", bold=True)
add_bullet("Defaults match AAOIFI; user can edit any number")

add_h2("4 operating principles")
add_bullet("Scholars hold religious authority. Mizan does not.")
add_bullet("Show the math.")
add_bullet("Credibility through citation.")
add_bullet("The strict view is a first-class profile.")

add_h2("Navigation cheatsheet")
add_bullet("/ — home   ·   /screen — screening tool   ·   /discover — trending + disagreement")
add_bullet("/methodology — citations   ·   /about — Miqdad as Founding Advisor   ·   /trust — honest scoreboard")
add_bullet("/custom — set your own   ·   /watchlist + /portfolio — saved tickers + compliance %")
add_bullet("/zakat   ·   /purification   ·   /ask — 49 Q&A   ·   /crypto   ·   /sukuk")

page_break()

# ═════════════════════════════════════════════════════════════
# THE SCRIPT — 0:00 OPENING
# ═════════════════════════════════════════════════════════════
add_h1("THE SCRIPT")

add_h2("0:00 – 2:00   OPENING", color=GOLD)
add_stage_direction("Wait for him to join. Smile when he does. Don't dive in immediately — let him settle.")

add_para("Greeting + thanks", bold=True, size=11)
add_script("Salaam Ibrahim — really good to meet you. Thanks so much for making the time.")
add_stage_direction("Pause. Let him respond.")

add_para("Brief small talk (max 30 sec)", bold=True, size=11)
add_script("How's the week treating you? / How's CUHK? / [whatever fits]")
add_stage_direction("Don't drag the small talk. 30 seconds, then transition.")

add_para("Frame the session", bold=True, size=11)
add_script("Quick framing before we dive in: this is 30 minutes. The format I'd love is — I'll share quickly what Mizan is, then you drive. You'll screen a stock you actually own or care about. I'll mostly stay quiet and watch. The whole point is I want to see your real experience — not me explaining things to you.")

add_script("I'm recording the session locally on my laptop so I can re-watch and catch things I miss in the moment. Is that OK with you?")
add_stage_direction("Wait for explicit YES. Start the recording. If he says no, recording stays off — handwrite notes only.")

add_para("Intro Mizan in 60 seconds", bold=True, size=11)
add_script("Let me share quickly what Mizan is, then I'll hand it over.")
add_script("Mizan is a halal investing screener. The short version: most halal investing apps — Zoya, Wahed, Musaffa — default everyone to AAOIFI standards and hide where scholars actually disagree. Mizan does the opposite. Six scholar profiles side-by-side. Every threshold cited to its primary source. We're the only screener I know of that actually implements Mufti Taqi Usmani's illiquid-asset rule — every other app skips it.")
add_script("I built it with my friend Miqdad as Founding Advisor — he's a Muslim, the reason this project exists, and he's gut-checked every methodology call.")
add_script("That's it. Now you drive. Have a stock in mind you'd like to screen, or want me to suggest one?")
add_stage_direction("Wait. He may say AAPL, TSLA, AMZN, NVDA, or something he owns. If he hesitates, suggest AAPL or AMZN. Then say: 'Great — share your screen, head to mizaninvesting.com, and just go.'")

add_h2("→ NOTES on opening")
add_note_lines(3)

page_break()

# ═════════════════════════════════════════════════════════════
# 2:00 - 18:00 WATCH PERIOD
# ═════════════════════════════════════════════════════════════
add_h1("2:00 – 18:00   WATCH PERIOD", shade=True)

add_warning("STAY QUIET. THIS IS THE CORE 16 MINUTES. Your job is to listen, watch, and write.")

add_h2("Rules during this period")
add_bullet("Let him drive. Don't narrate.", bold=True, color=RED)
add_bullet("Don't explain features when he's confused — WRITE IT DOWN", bold=True, color=RED)
add_bullet("Don't fix bugs / typos / weird UI in real time — WRITE IT DOWN", bold=True, color=RED)
add_bullet("Watch where his eye goes first on each page")
add_bullet("Watch what he clicks, hovers over, scrolls past")
add_bullet("Write his EXACT WORDS verbatim when he speaks")

add_h2("Sparingly — only when he pauses, use these prompts:")
add_script("Walk me through what you're doing — narrate it for me.")
add_script("Where did your eye go first?")
add_script("What did you expect that didn't happen?")
add_script("Talk me through what you're reading right now.")

add_h2("→ LIVE NOTES (write verbatim quotes in quotation marks)")
add_note_lines(8)

page_break()
add_note_lines(10)

page_break()

# ═════════════════════════════════════════════════════════════
# 18:00 - 26:00 DIAGNOSTIC QUESTIONS
# ═════════════════════════════════════════════════════════════
add_h1("18:00 – 26:00   DIAGNOSTIC QUESTIONS", shade=True)

add_stage_direction("Transition gently when he's done exploring. Don't rush him — if he's still actively using it, give another 2-3 min.")

add_para("Transition line:", bold=True, size=11)
add_script("OK — let me ask you a few specific things. Be brutally honest, please. 'Looks cool' is the least useful thing you can say.")

add_h2("Q1 — One-sentence description")
add_script("If you described Mizan to one of your friends in ONE sentence, what would you say? Don't polish it.")
add_note_lines(3, label="His ONE-SENTENCE description (verbatim)")

add_h2("Q2 — The one thing")
add_script("What's the ONE thing you'd change first?")
add_note_lines(3, label="The ONE thing")

add_h2("Q3 — Did the verdict match?")
add_script("Did the verdict on the stock you screened match what you expected? Why or why not?")
add_note_lines(4, label="Verdict match + reasoning")

add_h2("Q4 — Scholar conflict")
add_script("If your scholar or imam disagreed with what Mizan said, would you change Mizan's verdict, or change scholars?")
add_note_lines(3, label="His answer")

add_h2("Q5 — The welcome banner")
add_script("What did you think when you first saw the welcome banner about AAOIFI being the default?")
add_stage_direction("This tests the new UX Miqdad designed. Critical to hear honest reaction.")
add_note_lines(3, label="Reaction to AAOIFI default banner")

add_h2("Q6 — Scholar picker")
add_script("Was anything confusing about picking a scholar?")
add_note_lines(3, label="Scholar-picker friction")

page_break()

# ═════════════════════════════════════════════════════════════
# 26:00 - 30:00 CLOSING
# ═════════════════════════════════════════════════════════════
add_h1("26:00 – 30:00   CLOSING", shade=True)

add_para("Transition", bold=True, size=11)
add_script("Last few quick ones before we wrap.")

add_h2("Closing Q1 — Intros")
add_script("Who else do you know — at CUHK MSA or elsewhere — who'd actually use this and tell me what's wrong? Can you intro me?")
add_note_lines(4, label="Names of people he'd intro (first names + context)")

add_h2("Closing Q2 — Quote permission")
add_script("OK if I quote anything you said publicly? Named, anonymous, or not at all?")
add_note_lines(2, label="Permission: NAMED / ANONYMOUS / NO / ASK FIRST → circle one")

add_h2("Closing Q3 — Repeat use")
add_script("What would make you come back and use Mizan every time you screen a stock?")
add_note_lines(4, label="Stickiness driver")

add_para("Sign-off", bold=True, size=11)
add_script("This has been incredibly helpful — thank you. I'll send you whatever I ship from your feedback within a week or two. Anything you want to add before we wrap?")
add_stage_direction("Pause. Wait. Sometimes the best line comes here, unprompted.")

add_script("Thanks again Ibrahim. JazakAllah khair for the time.")
add_stage_direction("End the call. STOP THE RECORDING. Save as 'ibrahim_malik_2026-06-04.mp4'.")

page_break()

# ═════════════════════════════════════════════════════════════
# QUESTIONS HE MIGHT ASK — ANSWERS READY
# ═════════════════════════════════════════════════════════════
add_h1("QUESTIONS HE MIGHT ASK — answers READY")
add_stage_direction("If he asks any of these mid-session or at the end, deliver these answers concisely. Don't ramble.")

add_h2("Q: How do you make money?")
add_script("Not making money yet — it's free. Long-term plan: a small subscription for portfolio tracking and advanced features, roughly five to ten dollars a month equivalent. Methodology stays free forever. No ads, no affiliate links.")

add_h2("Q: Where's the financial data from?")
add_script("Financial Modeling Prep — institutional-grade equity data provider. Updated quarterly. I'm transparent about the data source on /methodology.")

add_h2("Q: Who validates the methodology? Are you a scholar?")
add_script("Not a scholar. I'm a CUHK med student. Miqdad is Founding Advisor — he's a Muslim and gut-checks methodology. Credibility right now comes from the citation chain — every threshold traces to AAOIFI Standard 21, or Mufti Taqi's published works, or DJIM's PDF, or Sistani's Code of Practice. Eventually we'll have a formal Shariah Supervisory Board. Until then, the honesty is the credibility. We don't claim authority we don't have.")

add_h2("Q: Why aren't you Muslim? Why should I trust this?")
add_script("Fair question. Honest answer: I'm not the source of authority — the cited scholars are. I defer to them on every religious call. The fact that I have no religious stake is actually why I can represent all six profiles fairly without favoring one. I show the math, cite every source, and never tell anyone what to follow.")

add_h2("Q: How is this different from Zoya / Musaffa / Wahed?")
add_script("Three things. First, multi-scholar — those default to AAOIFI, we show all six side-by-side. Second, Mufti Taqi's illiquid-asset rule — we're the only one that implements it; the others don't even check it. Third, every threshold cited — they give black-box scores, we show the math behind every verdict.")

add_h2("Q: Does it have [stock he owns]?")
add_script("Prototype dataset is around 163 representative tickers right now. Production version supports any listed equity — I'm adding more weekly. If yours isn't there, tell me and I'll add it in the next batch.")

add_h2("Q: Can I save my watchlist / track my portfolio?")
add_script("Yes — /watchlist for saved tickers, /portfolio to see your overall compliance percentage. Both work without an account — stored in your browser.")

add_h2("Q: What about crypto? Sukuk?")
add_script("Yes — /crypto for multi-scholar views on Bitcoin, ETH, etc.; /sukuk for Islamic bonds. Both are in earlier development than the stock screener.")

add_h2("Q: Can I use this to actually invest?")
add_script("Mizan provides INFORMATION, not investment advice. The verdicts are scholarly compliance opinions, not buy or sell recommendations. You'd still execute through a regular brokerage — Mizan tells you which stocks pass your scholar's bar, and you decide what to do with that.")

add_h2("Q: Why CUHK med student building this — what's your background in finance?")
add_script("None formally. Software builder, careful researcher, can read a balance sheet. The actual halal screening logic is just disciplined arithmetic on published thresholds. I'm not generating new rulings — I'm faithfully implementing existing ones with full citation.")

page_break()

# ═════════════════════════════════════════════════════════════
# CRITICAL RULES
# ═════════════════════════════════════════════════════════════
add_h1("CRITICAL SESSION RULES")

add_h2("DO", color=GREEN)
add_bullet("Stay quiet during the use period. Silence is data.")
add_bullet("Write his EXACT words verbatim. Not summaries.")
add_bullet("Acknowledge feedback with 'Interesting' or 'Tell me more' — neutral, not defensive.")
add_bullet("Be honest if you don't know something: 'I don't know — let me find out and email you.'")
add_bullet("Pause before responding. A two-second pause makes you look thoughtful.")
add_bullet("Smile occasionally. He should feel comfortable being honest.")

add_h2("DON'T", color=RED)
add_warning("DON'T fix bugs in real time. WRITE THEM DOWN.")
add_warning("DON'T explain features when he's confused. WATCH the confusion.")
add_warning("DON'T promise features. Say 'Interesting, I'll think about it' — never 'yes I'll add that.'")
add_warning("DON'T defend Mizan when criticized. Listen and write.")
add_warning("DON'T lead him: 'Did you find this confusing?' is leading. 'What do you think?' is not.")
add_warning("DON'T pitch Mizan during the session. He's already there.")
add_warning("DON'T over-explain Miqdad's role. If asked, brief: 'He's the reason this exists, and he gut-checks methodology.'")

# ═════════════════════════════════════════════════════════════
# CUSHION NOTES PAGE
# ═════════════════════════════════════════════════════════════
page_break()
add_h1("OVERFLOW NOTES")
add_stage_direction("If you run out of space in the boxes above, write here. Date stamp anything important.")
add_note_lines(20)

page_break()
add_note_lines(20)

# ═════════════════════════════════════════════════════════════
# POST-CALL CHECKLIST
# ═════════════════════════════════════════════════════════════
page_break()
add_h1("POST-CALL CHECKLIST — do within 60 min of hanging up")

add_bullet("Save Zoom recording: 'ibrahim_malik_2026-06-04.mp4'")
add_bullet("Type 3-5 verbatim quotes into the tracker → Quotable Feedback sheet")
add_bullet("Update Outreach Log → Ibrahim row → outcome column: 'Session done. Top 3 things: [list].'")
add_bullet("List bugs / UX issues found in a new section of OUTREACH.md or a tickets.md")
add_bullet("Send Ibrahim a 2-line thank-you email with top 3 things you're working on from his feedback")
add_bullet("If he gave intro names — draft warm-intro DMs for each, ready Friday")
add_bullet("Update Day 7 row in Daily Counts: +1 quotable, +1 community engagement, notes")
add_bullet("Ping Claude with the verbatim quotes — we'll start drafting your first public LinkedIn post")

add_hr()

add_para("Document built for Gordon Louie · Mizan founder · Hong Kong / CUHK med school",
         size=9, color=GRAY, italic=True)
add_para("This session is the single highest-leverage 30 minutes of Mizan's first 30 days. Don't wing it.",
         size=9, color=GOLD, italic=True)

# ═════════════════════════════════════════════════════════════
# SAVE
# ═════════════════════════════════════════════════════════════
output = "/sessions/busy-sleepy-faraday/mnt/mizan/Ibrahim_Session_Script.docx"
doc.save(output)
print(f"Saved: {output}")
