"""Build Mizan Outreach Tracker (.xlsx for Apple Numbers / Pages)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------- Style helpers ----------
NAVY = "0B1F3A"
GOLD = "C9A961"
CREAM = "F4ECD8"
LIGHT = "F5F5F5"

font_title = Font(name="Arial", size=18, bold=True, color="FFFFFF")
font_subtitle = Font(name="Arial", size=10, italic=True, color="888888")
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
font_body = Font(name="Arial", size=11, color="000000")
font_small = Font(name="Arial", size=9, color="555555")

fill_navy = PatternFill("solid", start_color=NAVY)
fill_gold = PatternFill("solid", start_color=GOLD)
fill_cream = PatternFill("solid", start_color=CREAM)
fill_light = PatternFill("solid", start_color=LIGHT)

thin = Side(border_style="thin", color="E5E5E5")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_header_row(ws, row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = center
        c.border = border


def style_title(ws, row, title, subtitle, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.font = font_title
    c.fill = fill_navy
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 32

    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
    sc = ws.cell(row=row + 1, column=1, value=subtitle)
    sc.font = font_subtitle
    sc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row + 1].height = 18


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- Workbook ----------
wb = Workbook()


# ===== Sheet 1: Outreach Log =====
ws = wb.active
ws.title = "Outreach Log"

style_title(ws, 1, "Outreach Log", "Every DM / connection request / email / IRL pitch. One row per send.", 14)

headers = [
    "Date sent", "Name", "Channel", "Template", "Personal context (1 line)",
    "Accepted?", "Replied?", "Reply date", "Quotable line (verbatim)",
    "Outcome", "Screened a stock?", "Shared unprompted?", "Next action", "Notes"
]
for col, h in enumerate(headers, start=1):
    ws.cell(row=4, column=col, value=h)
style_header_row(ws, 4, len(headers))
ws.row_dimensions[4].height = 38

# Example rows (greyed-out, instructive)
example_rows = [
    ["2026-05-29", "Mufti Faraz Adam", "LinkedIn connect", "Personalised", "UK Hanafi scholar, Amanah Advisors",
     "—", "—", "", "", "Pending", "—", "—", "Wait 48h then engage post", "Tier 1 named target"],
    ["2026-05-29", "Ibrahim Khan", "LinkedIn connect", "Personalised", "IFG co-founder",
     "—", "—", "", "", "Pending", "—", "—", "Wait 48h then engage post", "Tier 1 named target"],
    ["", "", "", "", "", "", "", "", "", "", "", "", "", ""],
]
for r_idx, row in enumerate(example_rows, start=5):
    for c_idx, val in enumerate(row, start=1):
        c = ws.cell(row=r_idx, column=c_idx, value=val)
        c.font = font_body if val and r_idx < 7 else font_small
        c.alignment = left
        c.border = border
        if r_idx < 7:
            c.fill = fill_cream

# Add 100 empty rows for future entries
for r in range(7, 107):
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=r, column=col)
        c.alignment = left
        c.border = border
        c.font = font_body

set_widths(ws, [12, 22, 18, 14, 32, 11, 11, 12, 42, 14, 13, 14, 28, 28])
ws.freeze_panes = "A5"


# ===== Sheet 2: Daily Counts =====
ws2 = wb.create_sheet("Daily Counts")
style_title(ws2, 1, "Daily Counts · 30-day exit criterion", "Pre-committed: <3 unprompted shares in 30 days = stop. Mark Day 1 as your first DM date.", 7)

# Summary
ws2.cell(row=4, column=1, value="Day 1 date:").font = Font(name="Arial", size=11, bold=True)
ws2.cell(row=4, column=2, value="").fill = fill_gold

ws2.cell(row=5, column=1, value="Exit date (Day 30):").font = Font(name="Arial", size=11, bold=True)
ws2.cell(row=5, column=2, value='=IF(ISBLANK(B4),"",B4+29)')
ws2.cell(row=5, column=2).fill = fill_gold

# Totals row
ws2.cell(row=7, column=1, value="Running totals →").font = Font(name="Arial", size=11, bold=True, italic=True)
ws2.cell(row=7, column=2, value='=SUM(B10:B40)').font = Font(name="Arial", size=11, bold=True)
ws2.cell(row=7, column=3, value='=SUM(C10:C40)').font = Font(name="Arial", size=11, bold=True)
ws2.cell(row=7, column=4, value='=SUM(D10:D40)').font = Font(name="Arial", size=11, bold=True)
ws2.cell(row=7, column=5, value='=SUM(E10:E40)').font = Font(name="Arial", size=11, bold=True)

for col in range(2, 6):
    ws2.cell(row=7, column=col).fill = fill_cream

ws2.cell(row=8, column=1, value="EXIT CRITERION CHECK").font = Font(name="Arial", size=10, bold=True, color="C9A961")
ws2.cell(row=8, column=5, value='=IF(E7>=3,"PROCEED — signal met","INSUFFICIENT — re-evaluate at day 30")')
ws2.cell(row=8, column=5).font = Font(name="Arial", size=10, bold=True)
ws2.cell(row=8, column=5).fill = fill_gold
ws2.merge_cells(start_row=8, start_column=5, end_row=8, end_column=7)

headers2 = ["Date", "DMs sent", "Replies received", "Quotable reactions",
            "Unprompted shares", "Communities engaged", "Notes"]
for col, h in enumerate(headers2, start=1):
    ws2.cell(row=9, column=col, value=h)
style_header_row(ws2, 9, len(headers2))
ws2.row_dimensions[9].height = 38

for r in range(10, 41):
    day_num = r - 9
    ws2.cell(row=r, column=1, value=f"Day {day_num}").font = font_small
    for col in range(2, 8):
        c = ws2.cell(row=r, column=col)
        c.alignment = left
        c.border = border
        c.font = font_body

set_widths(ws2, [16, 12, 16, 16, 16, 16, 32])


# ===== Sheet 3: Quotable Feedback =====
ws3 = wb.create_sheet("Quotable Feedback")
style_title(ws3, 1, "Quotable Feedback", "Verbatim quotes only. These anchor your first public LinkedIn post.", 6)

headers3 = ["Date", "Name", "Exact quote (verbatim)", "Source / context",
            "OK to use publicly?", "Used where"]
for col, h in enumerate(headers3, start=1):
    ws3.cell(row=4, column=col, value=h)
style_header_row(ws3, 4, len(headers3))
ws3.row_dimensions[4].height = 38

# Pre-seed the two reactions you already have
seed_quotes = [
    ["2026-05-26", "Tester 1", "this is a cool site", "Reaction over dinner", "No (deflected)", ""],
    ["2026-05-27", "Tester 2", "transparency builds trust immediately. add more real examples.",
     "WhatsApp after walk-through", "Ask first", ""],
]
for r_idx, row in enumerate(seed_quotes, start=5):
    for c_idx, val in enumerate(row, start=1):
        c = ws3.cell(row=r_idx, column=c_idx, value=val)
        c.font = font_body
        c.alignment = left
        c.border = border
        c.fill = fill_cream

for r in range(7, 50):
    for col in range(1, len(headers3) + 1):
        c = ws3.cell(row=r, column=col)
        c.alignment = left
        c.border = border
        c.font = font_body

set_widths(ws3, [12, 22, 60, 28, 18, 22])
ws3.freeze_panes = "A5"


# ===== Sheet 4: Hanafi Communities Log =====
ws4 = wb.create_sheet("Hanafi Communities")
style_title(ws4, 1, "Hanafi Communities", "Lurk first, comment substantively, never pitch on first action.", 6)

headers4 = ["Community", "URL / handle", "Status",
            "First action date", "First action (what you did)", "Notes"]
for col, h in enumerate(headers4, start=1):
    ws4.cell(row=4, column=col, value=h)
style_header_row(ws4, 4, len(headers4))
ws4.row_dimensions[4].height = 38

seed_communities = [
    ["IFG Discord", "islamicfinanceguru.com/discord", "Not yet joined", "", "", "Tier 1 — Ibrahim Khan + Faraz Adam built this. Hanafi-heavy."],
    ["@MuftiFarazAdam (X)", "x.com/MuftiFarazAdam", "Not engaged", "", "", "Tier 1 — UK Hanafi scholar."],
    ["Practical Islamic Finance (YouTube)", "youtube.com/@PracticalIslamicFinance", "Not engaged", "", "", "Tier 1 — Rakaan Kayali, Hanafi-leaning."],
    ["r/IslamicFinance (Reddit)", "reddit.com/r/IslamicFinance", "Not engaged", "", "", "Tier 2 — broader Muslim finance audience."],
    ["East London Mosque WhatsApp", "—", "Not joined", "", "", "Tier 2 — UK Hanafi community."],
    ["Deoband / Tablighi Telegram", "—", "Not joined", "", "", "Tier 2 — strict Hanafi, Mufti Taqi resonates."],
]
for r_idx, row in enumerate(seed_communities, start=5):
    for c_idx, val in enumerate(row, start=1):
        c = ws4.cell(row=r_idx, column=c_idx, value=val)
        c.font = font_body
        c.alignment = left
        c.border = border

for r in range(11, 25):
    for col in range(1, len(headers4) + 1):
        c = ws4.cell(row=r, column=col)
        c.alignment = left
        c.border = border
        c.font = font_body

set_widths(ws4, [32, 38, 18, 16, 32, 38])
ws4.freeze_panes = "A5"


# ===== Sheet 5: Target List =====
ws5 = wb.create_sheet("Target Queue")
style_title(ws5, 1, "Target Queue", "Pre-load names here. Move to Outreach Log once sent.", 6)

headers5 = ["Tier", "Name", "LinkedIn URL / handle", "Channel",
            "Template", "Personal hook (one line)"]
for col, h in enumerate(headers5, start=1):
    ws5.cell(row=4, column=col, value=h)
style_header_row(ws5, 4, len(headers5))
ws5.row_dimensions[4].height = 38

seed_targets = [
    ["1", "Mufti Faraz Adam", "linkedin.com/in/muftifarazadam", "LinkedIn connect", "Custom-1", "UK Hanafi scholar, Amanah Advisors"],
    ["1", "Ibrahim Khan", "linkedin.com/in/ibrahim-khan-ifg", "LinkedIn connect", "Custom-2", "IFG co-founder, halal investing thought leader"],
    ["1", "Mohsin Patel", "linkedin.com/in/mohsinpatel", "LinkedIn connect", "Custom-3", "IFG co-founder"],
    ["1", "Joe Bradford", "linkedin.com/in/joebradfordnet", "LinkedIn connect", "Custom-4", "US Islamic scholar, finance ethics"],
    ["1", "Rakaan Kayali", "linkedin.com/in/rakaankayali", "LinkedIn connect", "Custom-5", "Practical Islamic Finance YouTube"],
    ["2", "(anyone at Wahed)", "linkedin search: Wahed", "LinkedIn connect", "Generic", "PM/eng/analyst at competitor"],
    ["2", "(anyone at Zoya)", "linkedin search: Zoya", "LinkedIn connect", "Generic", ""],
    ["2", "(anyone at Musaffa)", "linkedin search: Musaffa", "LinkedIn connect", "Generic", ""],
    ["3", "(CUHK Muslim alumni in finance)", "linkedin search", "LinkedIn connect", "Generic", "HK local network"],
]
for r_idx, row in enumerate(seed_targets, start=5):
    for c_idx, val in enumerate(row, start=1):
        c = ws5.cell(row=r_idx, column=c_idx, value=val)
        c.font = font_body
        c.alignment = left
        c.border = border
        tier = row[0]
        if tier == "1":
            c.fill = fill_cream

for r in range(14, 40):
    for col in range(1, len(headers5) + 1):
        c = ws5.cell(row=r, column=col)
        c.alignment = left
        c.border = border
        c.font = font_body

set_widths(ws5, [8, 36, 36, 16, 14, 40])
ws5.freeze_panes = "A5"


# ===== Sheet 6: Templates reference =====
ws6 = wb.create_sheet("Templates")
style_title(ws6, 1, "DM Templates", "Quick reference. Full versions in OUTREACH.md.", 2)

templates = [
    ("Template A — Hanafi / Mufti Taqi followers (strongest cold lead)",
     "Hey [name] — quick one. Are you (or anyone you know) Hanafi-leaning, or following Mufti Taqi Usmani's framework on investing? I built mizaninvesting.com — first halal screener that actually implements his illiquid-asset rule. Every other app ignores it. Would love 5 mins of honest feedback if you've ever tried Zoya or Musaffa and walked away frustrated. No pitch, just curious what you'd find broken."),
    ("Template B — Finance-curious Muslim (any school)",
     "Hey [name] — building a halal investing screener (mizaninvesting.com) and you're one of the few people I know who actually thinks about this stuff carefully. Would you screen one stock you actually own and tell me what felt wrong? Not 'is the site cool' — what's wrong with the verdict or the math. 5 mins."),
    ("Template C — Non-Muslim finance/tech (informational, no ask)",
     "Hey [name] — small update: I shipped mizaninvesting.com, a halal investing screener with six scholarly profiles and fully-cited methodology. Sharing in case you know any Muslim investors who'd find it useful. No favour needed."),
    ("Custom-1 — Mufti Faraz Adam (LinkedIn connect note, 300ch)",
     "Asalaamu alaikum Mufti Faraz — Gordon, CUHK medical student in Hong Kong. I built mizaninvesting.com, a multi-scholar halal screener. The only one I know of that implements Mufti Taqi's illiquid-asset rule properly. Would deeply value your eye on the methodology page if you ever have a moment."),
    ("Custom-2 — Ibrahim Khan",
     "Hi Ibrahim — Gordon, CUHK med student building Mizan (mizaninvesting.com), a halal screener with six scholarly profiles and full primary-source citations. I've been reading IFG content for two years to understand the landscape. Would value the connection, no ask attached."),
    ("Custom-3 — Mohsin Patel",
     "Hi Mohsin — Gordon from Hong Kong, CUHK med student. Shipped mizaninvesting.com — multi-scholar halal screener, every threshold cited to a primary source. IFG's 'is Apple halal' series was part of what showed me the gap. Would value the connection."),
    ("Custom-4 — Joe Bradford",
     "Asalaamu alaikum Sheikh Joe — Gordon, CUHK medical student. I built a halal investing screener (mizaninvesting.com) that cites every threshold to its primary source — AAOIFI 21, Mufti Taqi's published works, DJIM methodology PDF. Would value the connection and any challenges to the methodology."),
    ("Custom-5 — Rakaan Kayali",
     "Hi Rakaan — Gordon from HK. I built mizaninvesting.com (multi-scholar halal screener, full citations on every threshold). Your stock-screening videos were one of the references I checked my work against. Would value the connection."),
    ("Generic LinkedIn connect (broader Muslim finance)",
     "Hi [name] — Gordon, CUHK med student in HK. I built mizaninvesting.com, a multi-scholar halal screener with full primary-source citations. Saw you [work at X / post about Islamic finance]. Would value the connection — no pitch, just building the network in this space."),
]
ws6.cell(row=4, column=1, value="Template").font = font_header
ws6.cell(row=4, column=1).fill = fill_navy
ws6.cell(row=4, column=2, value="Body").font = font_header
ws6.cell(row=4, column=2).fill = fill_navy
ws6.cell(row=4, column=1).alignment = center
ws6.cell(row=4, column=2).alignment = center
ws6.row_dimensions[4].height = 28

for r_idx, (name, body) in enumerate(templates, start=5):
    c1 = ws6.cell(row=r_idx, column=1, value=name)
    c2 = ws6.cell(row=r_idx, column=2, value=body)
    c1.font = Font(name="Arial", size=10, bold=True)
    c2.font = font_body
    c1.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c1.border = border
    c2.border = border
    ws6.row_dimensions[r_idx].height = 100

set_widths(ws6, [42, 92])


# ---------- Save ----------
output = "/sessions/busy-sleepy-faraday/mnt/mizan/Mizan_Outreach_Tracker.xlsx"
wb.save(output)
print(f"Saved: {output}")
